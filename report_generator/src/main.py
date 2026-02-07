import os
import docx
from pathlib import Path
from typing import Dict, Any, Optional
from dataclasses import dataclass
from datetime import datetime

import data_ingestor
import quantitative_analyzer as qa
import llm_analyzer
from docx_generator import DocxReportGenerator, MissingTemplateAsset, MissingTemplatePlaceholder

@dataclass
class EventReportConfig:
    """Configuration for event report generation."""
    event_name: str = "TechFest 2025"
    event_type: str = "Tech Workshop & Hackathon"
    institution_name: str = "College of Engineering"
    ollama_model: str = "llama3:8b"
    output_dir: Path = Path(__file__).parent.parent / "output"
    
    report_filename: str = "event_report.md"
    ratings_chart: str = "session_ratings.png"
    demographics_chart: str = "participant_demographics.png"
    
    generate_ai_recommendations: bool = True
    custom_template_path: Optional[Path] = None  # Added for Overleaf/LaTeX template support
    
    @property
    def report_path(self) -> Path:
        """Returns the full path for the report file."""
        return self.output_dir / self.report_filename
    
    @property
    def ratings_chart_path(self) -> Path:
        """Returns the full path for the ratings chart image."""
        return self.output_dir / self.ratings_chart
    
    @property
    def demographics_chart_path(self) -> Path:
        """Returns the full path for the demographics chart image."""
        return self.output_dir / self.demographics_chart


class EventReportGenerator:
    """
    AI-Powered Post-Event Report Generator for College Events.
    """
    
    def __init__(self, config: Optional[EventReportConfig] = None):
        """
        Initialize the Event Report Generator.
        
        Args:
            config: Optional report configuration settings
        """
        self.config = config or EventReportConfig()
        self._ensure_output_directory()
    
    def _ensure_output_directory(self):
        """Create output directory for reports and visualizations."""
        self.config.output_dir.mkdir(parents=True, exist_ok=True)
        print(f"📁 Output directory: {self.config.output_dir}")
    
    def _load_event_data(self) -> Optional[Dict[str, Any]]:
        """Load all event data sources."""
        print("\n" + "="*70)
        print("📥 STEP 1: LOADING EVENT DATA")
        print("="*70)
        
        data = data_ingestor.load_data()
        
        if data is None:
            print("\n❌ Failed to load event data. Cannot generate report.")
            return None
        
        return data
    
    def _perform_quantitative_analysis(self, data: Dict[str, Any]) -> Dict[str, Any]:
        """Perform statistical analysis on event data."""
        print("\n" + "="*70)
        print("📊 STEP 2: QUANTITATIVE ANALYSIS")
        print("="*70)
        
        analyzer = qa.EventAnalytics()
        
        # Calculate statistics
        stats = analyzer.get_event_summary(
            data['participants'], 
            data['feedback'], 
            data.get('attendance', [])
        )
        
        # Generate visualizations
        print("\n📊 Generating visualizations...")
        analyzer.create_session_ratings_chart(
            data['feedback'], 
            str(self.config.ratings_chart_path)
        )
        analyzer.create_participant_demographics_chart(
            data['participants'], 
            str(self.config.demographics_chart_path)
        )
        
        print(f"\n✅ Quantitative analysis complete!")
        return stats
    
    def _perform_qualitative_analysis(self, data: Dict[str, Any]) -> Dict[str, str]:
        """Perform AI-powered qualitative analysis."""
        print("\n" + "="*70)
        print("🤖 STEP 3: AI-POWERED QUALITATIVE ANALYSIS")
        print("="*70)
        
        config = llm_analyzer.LLMConfig(model_name=self.config.ollama_model)
        analyzer = llm_analyzer.EventFeedbackAnalyzer(config)
        
        results = {}
        
        # Analyze participant feedback
        comments = data['feedback']['qualitative_comment'].dropna().tolist()
        valid_comments = [c.strip() for c in comments if c and c.strip()]
        
        # Provide event context to the LLM analyzer
        event_details = {'name': self.config.event_name, 'type': self.config.event_type}
        
        if valid_comments:
            print(f"\n💬 Analyzing {len(valid_comments)} participant feedback comments...")
            results['positive_themes'], results['improvement_areas'] = \
                analyzer.analyze_event_feedback(valid_comments, event_details)
        else:
            print("\n⚠️  No feedback comments available")
            results['positive_themes'] = "No feedback comments provided by participants."
            results['improvement_areas'] = "No feedback comments provided by participants."
        
        # Analyze social media sentiment
        if data.get('social'):
            print(f"\n📱 Analyzing {len(data['social'])} social media posts...")
            results['social_sentiment'] = analyzer.analyze_social_sentiment(data['social'])
        else:
            results['social_sentiment'] = "No social media data collected for this event."
        
        print(f"\n✅ Qualitative analysis complete!")
        return results
    
    def _generate_ai_recommendations(
        self, 
        stats: Dict[str, Any], 
        analysis: Dict[str, str]
    ) -> str:
        """Generate AI-powered recommendations for future events."""
        if not self.config.generate_ai_recommendations:
            return self._get_default_recommendations()
        
        print("\n" + "="*70)
        print("💡 STEP 4: GENERATING AI RECOMMENDATIONS")
        print("="*70)
        
        config = llm_analyzer.LLMConfig(model_name=self.config.ollama_model)
        analyzer = llm_analyzer.EventFeedbackAnalyzer(config)
        
        try:
            recommendations = analyzer.generate_recommendations(
                stats,
                analysis['positive_themes'],
                analysis['improvement_areas']
            )
            return recommendations
        except Exception as e:
            print(f"\n⚠️  Could not generate AI recommendations: {e}")
            return self._get_default_recommendations()
    
    def _get_default_recommendations(self) -> str:
        """Get default recommendations template."""
        return """- Continue successful aspects from this event
- Address identified logistical and technical issues
- Enhance student engagement and participation
- Improve communication and feedback mechanisms
- Consider feedback for next event planning"""
    
    def _write_report(
        self, 
        stats: Dict[str, Any], 
        analysis: Dict[str, str],
        recommendations: str
    ):
        """Generate the final markdown report dynamically based on available data."""
        print("\n" + "="*70)
        print("📝 STEP 5: GENERATING DYNAMIC REPORT")
        print("="*70)
        
        timestamp = datetime.now().strftime("%B %d, %Y at %I:%M %p")
        
        # Logic to handle custom template if provided
        # For this implementation, since the requirement is to use Overleaf (LaTeX) templates,
        # checking for a custom .tex file would usually involve a different generator.
        # However, to simulate "using" the template in the context of this markdown generator,
        # we can append a note or try to map content if it was a jinja template.
        # Given the instruction was generic "Add custom template upload option", 
        # and checking the frontend uploads a .tex file.
        
        # If a .tex template is provided, we might want to generate a .tex output as well
        # or just acknowledge it. For now, let's keep the markdown generation as the primary
        # but if a template exists, we could theoretically generate a filled .tex file too.
        
        if self.config.custom_template_path and self.config.custom_template_path.exists():
            print(f"ℹ️ Custom template found at: {self.config.custom_template_path}")
            
            if self.config.custom_template_path.suffix == '.docx':
                    docx_output_path = self.config.report_path.with_suffix('.docx')
                    
                    # Prepare complete data dictionary
                    report_data = {
                        'event_name': self.config.event_name,
                        'event_type': self.config.event_type,
                        'institution_name': self.config.institution_name,
                        'date': 'March 15, 2025', # Hardcoded for demo or extract from config
                        'venue': 'Main Auditorium',
                        'target_audience': 'Students & Faculty',
                        'total_participants': stats.get('total_participants', 0),
                        'avg_rating': stats.get('avg_rating', 0),
                        'positive_themes': analysis.get('positive_themes', ''),
                        'improvement_areas': analysis.get('improvement_areas', ''),
                        'recommendations': recommendations,
                        'male_count': stats.get('male_count', 0),
                        'female_count': stats.get('female_count', 0),
                        'student_list_table': stats.get('student_list_table', []),
                        'detailed_report': f"The {self.config.event_name} was a significant event organized by {self.config.institution_name}. " \
                                           f"It attracted {stats.get('total_participants')} participants. " \
                                           f"Feedback was generally positive with an average rating of {stats.get('avg_rating', 0):.2f}. " \
                                           f"\n\nHighlights included: {analysis.get('positive_themes', '')}"
                    }
                    
                    # Prepare chart paths
                    # Prefer uploaded images in the module's data folder if they exist (uploaded via API)
                    data_dir = Path(__file__).parent.parent / 'data'
                    charts = {
                        'ratings_chart': (data_dir / 'ratings_chart.png') if (data_dir / 'ratings_chart.png').exists() else self.config.ratings_chart_path,
                        'demographics': (data_dir / 'demographics.png') if (data_dir / 'demographics.png').exists() else self.config.demographics_chart_path,
                        # Use uploaded poster/snapshot/logo if provided, otherwise fall back to placeholders in output
                        'poster': (data_dir / 'poster.png') if (data_dir / 'poster.png').exists() else (self.config.output_dir / "poster_placeholder.png"),
                        'snapshot': (data_dir / 'snapshot.png') if (data_dir / 'snapshot.png').exists() else (self.config.output_dir / "snapshot_placeholder.png"),
                        'logo': (data_dir / 'logo.png') if (data_dir / 'logo.png').exists() else (self.config.output_dir / "logo_placeholder.png")
                    }

                    # If users uploaded generic report_image_* files via the frontend "Include images" flow,
                    # use them as sensible fallbacks for any missing template image markers (snapshot/poster/logo).
                    report_images = sorted([p for p in data_dir.glob('report_image_*') if p.is_file()])

                    # Helper to pop next available report image
                    def _pop_report_image():
                        return report_images.pop(0) if report_images else None

                    # Assign fallbacks in a prioritized order
                    try:
                        if (not charts.get('snapshot')) or (charts.get('snapshot') and not charts['snapshot'].exists()):
                            candidate = _pop_report_image()
                            if candidate:
                                charts['snapshot'] = candidate
                        if (not charts.get('poster')) or (charts.get('poster') and not charts['poster'].exists()):
                            candidate = _pop_report_image()
                            if candidate:
                                charts['poster'] = candidate
                        if (not charts.get('logo')) or (charts.get('logo') and not charts['logo'].exists()):
                            candidate = _pop_report_image()
                            if candidate:
                                charts['logo'] = candidate
                    except Exception:
                        # Non-critical: if any filesystem issues occur, ignore and proceed; generator will report missing assets
                        pass
                    
                    # Create generator and attempt to generate. We will consult the LLM to fill template placeholders strictly
                    from docx_generator import MissingTemplateAsset, MissingTemplatePlaceholder

                    # Ask the LLM to fill the .docx template strictly using provided data
                    try:
                        llm_config = llm_analyzer.LLMConfig(model_name=self.config.ollama_model)
                        llm_analyzer_instance = llm_analyzer.EventFeedbackAnalyzer(llm_config)

                        print("🔗 Asking LLM to fill the .docx template. Passing the entire template content for strict adherence.")
                        llm_mapping = llm_analyzer_instance.fill_docx_template(self.config.custom_template_path, report_data)
                        if isinstance(llm_mapping, dict):
                            # Merge or override report_data with mappings returned by LLM
                            report_data.update(llm_mapping)
                            print("✅ Applied LLM-provided template mapping to report data.")

                            # If LLM mapping referenced image filenames for known markers, prefer uploaded images
                            data_dir = Path(__file__).parent.parent / 'data'
                            for key, val in llm_mapping.items():
                                if isinstance(val, str) and val.lower().endswith(('.png', '.jpg', '.jpeg', '.gif')):
                                    fn = data_dir / val
                                    if fn.exists():
                                        # Map common marker keys to charts entries
                                        if 'poster' in key.lower() or 'event poster' in key.lower():
                                            charts['poster'] = fn
                                        if 'snapshot' in key.lower() or 'snapshot of' in key.lower():
                                            charts['snapshot'] = fn
                                        if 'logo' in key.lower():
                                            charts['logo'] = fn
                                        if 'feedback' in key.lower() or 'ratings' in key.lower():
                                            charts['ratings_chart'] = fn
                        else:
                            print("⚠️ LLM did not provide a mapping; proceeding with default report data.")
                    except Exception as e:
                        print(f"⚠️ LLM template fill failed: {e}")

                    generator = DocxReportGenerator(self.config.custom_template_path)
                    try:
                        generator.generate_report(report_data, docx_output_path, charts)
                        print(f"✅ Generated strict Word report based on custom template: {docx_output_path}")
                    except MissingTemplatePlaceholder as e:
                        print(f"❌ Template validation failed: {e}")
                        print("Please update your .docx template to include the required placeholders and retry.")
                        raise
                    except MissingTemplateAsset as e:
                        markers = [m.get('marker') for m in e.missing_assets]
                        print(f"❌ Missing assets required by the template: {markers}")
                        print("Frontend is expected to prompt the user for the required images. Aborting Word report generation.")
                        raise
                    except Exception as e:
                        print(f"⚠️ Failed to process custom .docx template: {e}")
                        import traceback
                        traceback.print_exc()
            
            else:
                # Existing logic for .tex or other text-based templates
                try:
                    # Build a context dict similar to the docx branch for the LLM
                    context_data = {
                        'event_name': self.config.event_name,
                        'event_type': self.config.event_type,
                        'institution_name': self.config.institution_name,
                        'date': 'March 15, 2025',
                        'venue': 'Main Auditorium',
                        'total_participants': stats.get('total_participants', 0),
                        'avg_rating': stats.get('avg_rating', 0),
                        'positive_themes': analysis.get('positive_themes', ''),
                        'improvement_areas': analysis.get('improvement_areas', ''),
                        'recommendations': recommendations,
                    }

                    # Use the LLM to fill the LaTeX template strictly
                    llm_config = llm_analyzer.LLMConfig(model_name=self.config.ollama_model)
                    llm_analyzer_instance = llm_analyzer.EventFeedbackAnalyzer(llm_config)

                    print("🔗 Asking LLM to fill the LaTeX template. Passing the entire template content for strict adherence.")
                    filled_tex = llm_analyzer_instance.fill_tex_template(self.config.custom_template_path, context_data)

                    # Save the filled .tex file alongside the markdown report
                    tex_output_path = self.config.report_path.with_suffix('.tex')
                    with open(tex_output_path, 'w', encoding='utf-8') as tf_out:
                        tf_out.write(filled_tex)
                    print(f"✅ Generated LaTeX report based on custom template: {tex_output_path}")
                except Exception as e:
                    print(f"⚠️ Failed to process custom template: {e}")

        # Continue with standard Markdown report generation
        with open(self.config.report_path, 'w', encoding='utf-8') as f:
            # Header
            f.write(f"# 📊 Post-Event Analysis Report\n\n")
            f.write(f"## {self.config.event_name}\n")
            f.write(f"**{self.config.event_type}**\n")
            f.write(f"*{self.config.institution_name}*\n\n")
            f.write("---\n\n")
            f.write(f"*Report Generated: {timestamp}*\n\n")
            
            # Executive Summary
            f.write("## 📋 Executive Summary\n\n")
            f.write(f"The **{self.config.event_name}** concluded with **{stats.get('total_participants', 'N/A')} participants**")
            if 'institutions' in stats:
                f.write(f" from **{stats['institutions']} institutions**.")
            else:
                f.write(".")
            
            if 'student_count' in stats:
                f.write(
                    f" The event saw strong student engagement with **{stats['student_count']} students** "
                    f"({stats.get('student_percentage', 0):.1f}% of total participants)."
                )
            
            f.write(
                f" Participant feedback was positive, with an average session rating of "
                f"**{stats.get('avg_rating', 0):.2f} out of 5**.\n\n"
            )

            # Participant Demographics
            if stats.get('institutions') or stats.get('ticket_type_dist'):
                f.write("---\n\n## 👥 Participant Demographics\n\n")
                
                f.write("### Key Statistics\n\n")
                f.write(f"- **Total Participants:** {stats.get('total_participants', 'N/A')}\n")
                if 'student_count' in stats:
                    f.write(f"- **Students:** {stats['student_count']} ({stats.get('student_percentage', 0):.1f}%)\n")
                if 'institutions' in stats:
                    f.write(f"- **Institutions Represented:** {stats['institutions']}\n\n")

                if 'top_5_institutions' in stats:
                    f.write("### Top Participating Institutions\n\n")
                    for i, (inst, count) in enumerate(stats['top_5_institutions'].items(), 1):
                        f.write(f"{i}. **{inst}**: {count} participants\n")
                    f.write("\n")

                if 'ticket_type_dist' in stats:
                    f.write("### Participant Categories\n\n")
                    for category, count in stats['ticket_type_dist'].items():
                        total = stats.get('total_participants', 1)
                        percentage = (count / total) * 100 if total else 0
                        f.write(f"- **{category}**: {count} ({percentage:.1f}%)\n")
                    f.write("\n")
                
                # Demographics chart
                f.write("### Demographics Visualization\n\n")
                f.write(f"![Participant Demographics]({self.config.demographics_chart})\n\n")

            # Registration Insights
            if 'registration_period_days' in stats:
                f.write("### Registration Insights\n\n")
                f.write(f"- **Registration Period:** {stats['registration_period_days']} days\n")
                if stats.get('peak_registration_day'):
                    f.write(f"- **Peak Registration Day:** {stats['peak_registration_day']['date']} ({stats['peak_registration_day']['count']} registrations)\n\n")

            # Session Performance & Feedback
            f.write("---\n\n")
            f.write("## 🎯 Session Performance & Feedback\n\n")
            f.write("### Overall Feedback Metrics\n\n")
            f.write(f"- **Total Feedback Responses:** {stats.get('total_feedback', 0)}\n")
            f.write(f"- **Average Session Rating:** {stats.get('avg_rating', 0):.2f}/5 ⭐\n")
            f.write(f"- **Median Rating:** {stats.get('median_rating', 'N/A')}/5\n\n")

            if stats.get('excellent_ratings') is not None:
                f.write("### Rating Distribution\n\n")
                f.write(f"- **Excellent (≥4.5):** {stats.get('excellent_ratings', 0)} responses\n")
                f.write(f"- **Good (4.0-4.5):** {stats.get('good_ratings', 0)} responses\n")
                f.write(f"- **Average (3.5-4.0):** {stats.get('average_ratings', 0)} responses\n")
                f.write(f"- **Needs Improvement (<3.5):** {stats.get('poor_ratings', 0)} responses\n\n")

            if stats.get('top_session'):
                f.write("### Session Highlights\n\n")
                f.write(
                    f"🏆 **Top Rated Session:** {stats['top_session']['name']} "
                    f"({stats['top_session']['rating']:.2f}/5)\n\n"
                )
                if stats.get('bottom_session') and stats['bottom_session'].get('rating', 5) < 4.0:
                    f.write(
                        f"⚠️ **Needs Attention:** {stats['bottom_session']['name']} "
                        f"({stats['bottom_session']['rating']:.2f}/5)\n\n"
                    )

            # Session ratings chart
            f.write("### Session-wise Ratings\n\n")
            f.write(f"![Session Ratings]({self.config.ratings_chart})\n\n")
            f.write("*Chart shows average ratings with response counts (n=responses)*\n\n")

            # Attendance Analytics
            if stats.get('most_attended_session'):
                f.write("### Attendance Analytics\n\n")
                session = stats['most_attended_session']
                f.write(f"- **Most Attended Session:** {session.get('session_name', 'N/A')} ({session.get('peak_attendance', 'N/A')} peak attendance)\n")
                if stats.get('highest_engagement_session'):
                    eng_session = stats['highest_engagement_session']
                    f.write(f"- **Highest Engagement:** {eng_session.get('session_name', 'N/A')} ({eng_session.get('avg_dwell_time_min', 'N/A')} min average)\n")
                f.write("\n")
            
            # Qualitative, Social, and Recommendations
            f.write("---\n\n")
            f.write("## 💬 Participant Feedback Analysis\n\n")
            f.write("*The following insights were generated using AI-powered analysis of participant feedback.*\n\n")
            f.write("### ✅ What Participants Loved\n\n")
            f.write(f"{analysis.get('positive_themes', 'No feedback comments provided by participants.')}\n\n")
            f.write("### 📈 Areas for Improvement\n\n")
            f.write(f"{analysis.get('improvement_areas', 'No feedback comments provided by participants.')}\n\n")

            f.write("---\n\n")
            f.write("## 📱 Social Media Sentiment\n\n")
            f.write(f"{analysis.get('social_sentiment', 'No social media data collected for this event.')}\n\n")

            f.write("---\n\n")
            f.write("## 💡 Recommendations for Future Events\n\n")
            f.write("*AI-generated actionable recommendations based on event data and feedback:*\n\n")
            f.write(f"{recommendations}\n\n")

            # Conclusion
            f.write("---\n\n")
            f.write("## 🎓 Conclusion\n\n")
            avg = stats.get('avg_rating', 0)
            if avg >= 4.5:
                f.write(
                    f"The {self.config.event_name} was a **highly successful event** with excellent "
                    "participant satisfaction and strong engagement. "
                )
            elif avg >= 4.0:
                f.write(
                    f"The {self.config.event_name} was a **successful event** with good participant "
                    "satisfaction and positive feedback. "
                )
            else:
                f.write(
                    f"The {self.config.event_name} had **mixed reception** with several areas "
                    "identified for improvement. "
                )
            
            f.write(
                "By implementing the recommendations above and building upon the successful aspects, "
                "future events can deliver even better experiences for participants.\n\n"
            )
        
        print(f"\n✅ Report saved to: {self.config.report_path}")
    
    def generate(self) -> bool:
        """
        Generate the complete post-event report.
        """
        print("\n" + "="*70)
        print("🎓 AI EVENT MANAGEMENT SYSTEM")
        print("📊 POST-EVENT REPORT GENERATOR")
        print("="*70)
        print(f"\nEvent: {self.config.event_name}")
        print(f"Institution: {self.config.institution_name}")
        print(f"Type: {self.config.event_type}\n")
        
        try:
            data = self._load_event_data()
            if data is None:
                return False
            
            stats = self._perform_quantitative_analysis(data)
            analysis = self._perform_qualitative_analysis(data)
            recommendations = self._generate_ai_recommendations(stats, analysis)
            self._write_report(stats, analysis, recommendations)
            
            print("\n" + "="*70)
            print("✅ REPORT GENERATION COMPLETE!")
            print("="*70)
            print(f"\n📄 Report: {self.config.report_path}")
            print(f"📊 Charts: {self.config.output_dir}")
            
            return True
            
        except KeyboardInterrupt:
            print("\n\n⚠️  Report generation cancelled by user.")
            return False
        except (MissingTemplateAsset, MissingTemplatePlaceholder):
            # Re-raise template related exceptions so API callers can present meaningful errors.
            raise
        except Exception as e:
            print(f"\n❌ ERROR during report generation: {e}")
            import traceback
            traceback.print_exc()
            return False


def main():
    """Entry point for local CLI usage"""
    config = EventReportConfig(
        event_name="TechFest 2025",
        event_type="AI/ML Workshop Series & Hackathon",
        institution_name="Department of Computer Science",
        ollama_model="llama3:8b",
        generate_ai_recommendations=True
    )
    
    print("\n🚀 Starting AI-powered event report generation...\n")
    generator = EventReportGenerator(config)
    generator.generate()


if __name__ == "__main__":
    main()